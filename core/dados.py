# -*- coding: utf-8 -*-
"""
core/dados.py

Lê a aba de dados brutos (histórico mensal por filial) do arquivo
"monitor_diretoria_....xlsm" que a diretoria da empresa gera todo mês,
e devolve uma lista de dicionários (um por filial) já filtrados pelo
mês/ano/unidade escolhidos.

Este módulo assume que o arquivo de origem tem uma aba (por padrão
chamada "Dados") com colunas, entre outras:
    Filial, Regiao, Faturamento, vlr-servico, valor-pecas,
    custo-pecas-medio, despesa-total, cst-estoque, Consultas,
    estoque, PMR, vlr-recebido, vlr-avencer, vlr-lp, lp-mes,
    vlr-atrasado, vlr-lp-recebido, desp-flex, mes, ano, unidade

Se a estrutura da planilha da diretoria mudar de nome de coluna no
futuro, ajuste apenas o dicionário RAW_FIELDS abaixo.
"""

import openpyxl

# Campos brutos que este programa precisa encontrar na aba de dados.
# A chave é o nome usado internamente pelo programa; o valor é o nome
# exato da coluna no arquivo da diretoria.
RAW_FIELDS = {
    'Filial': 'Filial',
    'Regiao': 'Regiao',
    'Faturamento': 'Faturamento',
    'vlr_servico': 'vlr-servico',
    'valor_pecas': 'valor-pecas',
    'custo_pecas_medio': 'custo-pecas-medio',
    'despesa_total': 'despesa-total',
    'cst_estoque': 'cst-estoque',
    'Consultas': 'Consultas',
    'estoque': 'estoque',
    'PMR': 'PMR',
    'vlr_recebido': 'vlr-recebido',
    'vlr_avencer': 'vlr-avencer',
    'vlr_lp': 'vlr-lp',
    'lp_mes': 'lp-mes',
    'vlr_atrasado': 'vlr-atrasado',
    'vlr_lp_recebido': 'vlr-lp-recebido',
    'desp_flex': 'desp-flex',
}


class ErroLeituraDados(Exception):
    pass


def carregar_dados_mensais(caminho_xlsm, mes, ano, unidade='Comércio', aba='Dados'):
    """
    Lê o arquivo mensal da diretoria e devolve uma lista de dicionários,
    um por filial, contendo os campos brutos necessários para gerar os
    relatórios.

    Parâmetros
    ----------
    caminho_xlsm : str
        Caminho completo do arquivo .xlsm/.xlsx mensal da diretoria.
    mes : int
        Mês a filtrar (1-12).
    ano : int
        Ano a filtrar (ex: 2026).
    unidade : str
        Unidade de negócio a filtrar (ex: 'Comércio' ou 'Indústria').
    aba : str
        Nome da aba dentro do arquivo que contém os dados brutos.

    Retorna
    -------
    list[dict]
        Um dicionário por filial, com as chaves internas de RAW_FIELDS.
    """
    try:
        wb = openpyxl.load_workbook(caminho_xlsm, data_only=True, read_only=True)
    except Exception as e:
        raise ErroLeituraDados(f"Não consegui abrir o arquivo:\n{caminho_xlsm}\n\nDetalhe: {e}")

    if aba not in wb.sheetnames:
        wb.close()
        raise ErroLeituraDados(
            f"A aba '{aba}' não existe nesse arquivo.\n"
            f"Abas encontradas: {', '.join(wb.sheetnames)}"
        )

    ws = wb[aba]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = next(linhas)

    colunas_necessarias = list(RAW_FIELDS.values()) + ['mes', 'ano', 'unidade']
    faltando = [col for col in colunas_necessarias if col not in cabecalho]
    if faltando:
        wb.close()
        raise ErroLeituraDados(
            "Essa aba de dados não tem as colunas esperadas: "
            + ", ".join(faltando)
            + "\n\nSe a estrutura do arquivo da diretoria mudou, avise quem mantém "
            "o programa para atualizar o mapeamento de colunas em core/dados.py."
        )

    idx = {nome: i for i, nome in enumerate(cabecalho)}

    registros = []
    for linha in linhas:
        d_bruto = dict(zip(cabecalho, linha))
        if d_bruto.get('mes') != mes or d_bruto.get('ano') != ano:
            continue
        if unidade and d_bruto.get('unidade') != unidade:
            continue
        registro = {chave: d_bruto.get(coluna) for chave, coluna in RAW_FIELDS.items()}
        registros.append(registro)

    wb.close()
    return registros


def listar_meses_disponiveis(caminho_xlsm, unidade='Comércio', aba='Dados'):
    """Devolve os pares (ano, mes) disponíveis no arquivo, do mais recente
    para o mais antigo — útil para popular a tela de seleção."""
    wb = openpyxl.load_workbook(caminho_xlsm, data_only=True, read_only=True)
    ws = wb[aba]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = next(linhas)
    idx_mes = cabecalho.index('mes')
    idx_ano = cabecalho.index('ano')
    idx_unid = cabecalho.index('unidade') if 'unidade' in cabecalho else None
    pares = set()
    for linha in linhas:
        if idx_unid is not None and linha[idx_unid] != unidade:
            continue
        if linha[idx_mes] is not None and linha[idx_ano] is not None:
            pares.add((int(linha[idx_ano]), int(linha[idx_mes])))
    wb.close()
    return sorted(pares, reverse=True)
