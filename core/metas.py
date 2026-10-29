# -*- coding: utf-8 -*-
"""
core/metas.py

Lida com a planilha de METAS por filial — o documento que diz qual é a
média/meta esperada para cada loja. Como esse arquivo pode mudar de
formato (novas colunas, nomes diferentes, ordem diferente) ao longo do
tempo, este módulo NÃO assume nomes de coluna fixos: em vez disso, guarda
um "mapeamento" (qual coluna da planilha de metas corresponde a qual
métrica do relatório) em data/metas_mapeamento.json.

Fluxo de uso:
1. ler_cabecalhos(caminho) -> devolve os nomes das colunas encontradas no
   arquivo de metas, para popular a tela de mapeamento.
2. Usuário escolhe, na interface, qual coluna é a Filial e quais colunas
   correspondem a quais métricas (ex: "Meta Faturamento" -> "Vlr
   Faturamento"). Pode deixar em branco as métricas que não têm meta.
3. salvar_mapeamento(mapeamento) -> grava essa escolha localmente.
4. carregar_metas(caminho, mapeamento) -> lê os valores de meta de cada
   filial usando o mapeamento salvo.

Quando o arquivo de metas mudar de estrutura no futuro, basta repetir o
passo 2 na tela "Metas" do programa — não precisa mexer em código.
"""

import json
import os
import openpyxl

from .appdir import base_dir, NOME_PASTA_DADOS

DATA_DIR = os.path.join(base_dir(), NOME_PASTA_DADOS)
MAPEAMENTO_PATH = os.path.join(DATA_DIR, 'metas_mapeamento.json')
METAS_CACHE_PATH = os.path.join(DATA_DIR, 'metas_ultimo_arquivo.json')

# Métricas do relatório que fazem sentido comparar com uma meta.
METRICAS_COM_META = [
    "Vlr Faturamento", "%Margem", "%Despesa", "Vlr Consultas",
    "%Consultas Fat.", "%Lucro", "PMR", "% Inad/Vencer",
]


class ErroMetas(Exception):
    pass


def ler_cabecalhos(caminho_arquivo, aba=None):
    """Devolve a lista de cabeçalhos (linha 1) de uma planilha de metas,
    para o usuário escolher na tela de mapeamento."""
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True, read_only=True)
    ws = wb[aba] if aba else wb.active
    primeira_linha = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    wb.close()
    return [str(c) for c in primeira_linha if c is not None]


def listar_abas(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, read_only=True)
    abas = wb.sheetnames
    wb.close()
    return abas


def salvar_mapeamento(caminho_arquivo, aba, mapeamento):
    """mapeamento: dict {'coluna_filial': 'Loja', 'Vlr Faturamento': 'Meta Fat.', ...}
    Só entram no dict as métricas que o usuário de fato mapeou."""
    os.makedirs(DATA_DIR, exist_ok=True)
    config = {
        'caminho_arquivo': caminho_arquivo,
        'aba': aba,
        'mapeamento': mapeamento,
    }
    with open(MAPEAMENTO_PATH, 'w', encoding='utf-8') as f:
        json.dump(config, f, ensure_ascii=False, indent=2)


def carregar_config_mapeamento():
    """Devolve o último mapeamento salvo, ou None se ainda não configurado."""
    if not os.path.exists(MAPEAMENTO_PATH):
        return None
    with open(MAPEAMENTO_PATH, encoding='utf-8') as f:
        return json.load(f)


def carregar_metas(caminho_arquivo, aba, mapeamento):
    """Lê a planilha de metas usando o mapeamento de colunas e devolve
    um dict {filial: {metrica: valor_meta}}."""
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True, read_only=True)
    ws = wb[aba] if aba else wb.active
    linhas = list(ws.iter_rows(values_only=True))
    cabecalho = [str(c) if c is not None else '' for c in linhas[0]]

    col_filial = mapeamento.get('coluna_filial')
    if not col_filial or col_filial not in cabecalho:
        wb.close()
        raise ErroMetas(
            "O mapeamento de metas não tem uma coluna de Filial válida. "
            "Refaça o mapeamento na aba 'Metas'."
        )
    idx_filial = cabecalho.index(col_filial)

    idx_por_metrica = {}
    for metrica in METRICAS_COM_META:
        col_planilha = mapeamento.get(metrica)
        if col_planilha and col_planilha in cabecalho:
            idx_por_metrica[metrica] = cabecalho.index(col_planilha)

    metas = {}
    for linha in linhas[1:]:
        if linha[idx_filial] is None:
            continue
        filial = str(linha[idx_filial]).strip()
        metas[filial] = {}
        for metrica, idx in idx_por_metrica.items():
            valor = linha[idx]
            if valor is not None:
                metas[filial][metrica] = valor
    wb.close()
    return metas
